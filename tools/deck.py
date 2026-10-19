"""Fill a company's own .pptx template with new content without touching its design.

    python tools/deck.py map   TEMPLATE.pptx MAP.yaml
    python tools/deck.py fill  TEMPLATE.pptx CONTENT.yaml OUT.pptx
    python tools/deck.py check TEMPLATE.pptx MAP.yaml CONTENT.yaml OUT.pptx [REPORT.json]

map:   lists every box that holds text (inside groups and table cells too), with
       its sample text, font size and roughly how many characters fit. A person
       marks the boxes to leave alone (logo, page number) with keep: true.
fill:  copies the template, drops slides the content doesn't use, and swaps the
       text of each named box in place: the first run of each paragraph keeps its
       formatting and only its text changes. Nothing is redrawn.
check: flags text that won't fit its box, sample text left behind, boxes whose
       font, size or color changed, and content keys that match no box.

See docs/deck-templates.md.
"""

import copy
import json
import re
import sys
from pathlib import Path

import yaml
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

EMU_PER_PT = 12700
FAILING = {"overflow", "wraps", "not_filled", "leftover", "style_changed", "unknown_key"}
CHROME_NAMES = ("Logo word", "Footer", "Slide number")  # boxes repeated on every slide
BULLET = re.compile(r"^([•\-–▪●]\s+)")
FONT_FILES = {  # Arial or a metric match, for measuring; swap for the template's own fonts
    False: ["arial.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"],
    True: ["arialbd.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"],
}


def load_font(bold, size):
    for f in FONT_FILES[bool(bold)]:
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    raise SystemExit("no Arial-compatible font found; install Arial or Liberation Sans")


# ---------- reading boxes ----------

def boxes(slide):
    """Yield (key, shape, text_frame, cell) for every text-holding box on a slide."""
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # group
                yield from walk(sh.shapes)
            elif sh.has_table:
                for r, row in enumerate(sh.table.rows):
                    for c, cell in enumerate(row.cells):
                        yield f"{sh.shape_id}.r{r}c{c}", sh, cell.text_frame, cell
            elif sh.has_text_frame and sh.text_frame.text.strip():
                yield str(sh.shape_id), sh, sh.text_frame, None
    yield from walk(slide.shapes)


def first_run(tf):
    for p in tf.paragraphs:
        for r in p.runs:
            return r
    return None


def style(tf):
    r = first_run(tf)
    if r is None:
        return {}
    f = r.font
    color = None
    try:
        color = str(f.color.rgb) if f.color and f.color.type is not None else None
    except AttributeError:
        pass
    return {"font": f.name, "size": f.size.pt if f.size else None, "bold": f.bold, "color": color}


def box_size(shape):
    """Usable width and height in points, after the box's inner margins."""
    tf = shape.text_frame
    w = shape.width - (tf.margin_left or Emu(91440)) - (tf.margin_right or Emu(91440))
    h = shape.height - (tf.margin_top or Emu(45720)) - (tf.margin_bottom or Emu(45720))
    return w / EMU_PER_PT, h / EMU_PER_PT


def wrap_lines(text, width_pt, size, bold):
    font = load_font(bold, int(round(size)))
    lines = 0
    for para in text.split("\n"):
        line = ""
        n = 1
        for word in para.split():
            trial = f"{line} {word}".strip()
            if font.getlength(trial) <= width_pt or not line:
                line = trial
            else:
                n += 1
                line = word
            while font.getlength(line) > width_pt and len(line) > 1:  # a word too long to fit breaks
                cut = max(i for i in range(1, len(line) + 1) if font.getlength(line[:i]) <= width_pt or i == 1)
                n += 1
                line = line[cut:]
        lines += n
    return lines


def capacity(width_pt, height_pt, size):
    lines = max(1, int(height_pt // (size * 1.2)))
    per_line = max(1, int(width_pt // (size * 0.5)))
    return lines, lines * per_line


# ---------- map ----------

def cmd_map(template, out):
    prs = Presentation(template)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        entries = []
        for key, shape, tf, cell in boxes(slide):
            st = style(tf)
            e = {"key": key, "name": shape.name, "sample": tf.text}
            if cell is None and st.get("size"):
                w, h = box_size(shape)
                e["max_lines"], e["max_chars"] = capacity(w, h, st["size"])
            e["size"] = st.get("size")
            if shape.name in CHROME_NAMES:
                e["chrome"] = True
            entries.append(e)
        slides.append({"slide": i, "use": "", "boxes": entries})
    Path(out).write_text(yaml.safe_dump({"template": str(template), "slides": slides},
                                        sort_keys=False, allow_unicode=True, width=100))
    print(f"mapped {sum(len(s['boxes']) for s in slides)} boxes on {len(slides)} slides -> {out}")


# ---------- fill ----------

def set_paragraphs(tf, lines):
    """Replace a text frame's text, keeping each paragraph's first-run formatting."""
    paras = list(tf.paragraphs)
    bullet = ""
    r0 = first_run(tf)
    if r0 is not None:
        m = BULLET.match(r0.text)
        bullet = m.group(1) if m else ""
    for i, line in enumerate(lines):
        if i < len(paras):
            p = paras[i]
        else:
            new = copy.deepcopy(paras[-1]._p)
            paras[-1]._p.addnext(new)
            p = [q for q in tf.paragraphs if q._p is new][0]
            paras.append(p)
        runs = p.runs
        if not runs:
            continue
        runs[0].text = bullet + line if bullet and not BULLET.match(line) else line
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)


def as_lines(value):
    return value if isinstance(value, list) else [str(value)]


def cmd_fill(template, content_path, out):
    content = yaml.safe_load(Path(content_path).read_text())
    prs = Presentation(template)
    keep = {int(s["slide"]) for s in content["slides"]}
    by_slide = {int(s["slide"]): s.get("boxes", {}) for s in content["slides"]}
    chrome = content.get("chrome", {})
    order = [int(s["slide"]) for s in content["slides"]]
    for i, slide in enumerate(prs.slides, 1):
        values = by_slide.get(i, {})
        found = {}
        for key, shape, tf, cell in boxes(slide):
            found[key] = (shape, tf)
            if shape.name in chrome and i in keep:
                n = order.index(i) + 1
                set_paragraphs(tf, [str(chrome[shape.name]).replace("{n}", str(n))])
        for key, value in values.items():
            key = str(key)
            if key not in found:
                continue  # reported by check
            shape, tf = found[key]
            if value is None:  # delete the box
                el = shape._element
                el.getparent().remove(el)
            elif isinstance(value, bool):
                sys.exit(f"slide {i} box {key}: {value!r} was read as true/false; put it in quotes")
            else:
                set_paragraphs(tf, as_lines(value))
    # drop unused slides, then reorder to the content's order
    sld_ids = prs.slides._sldIdLst
    ids = list(sld_ids)
    for i, sid in enumerate(ids, 1):
        if i not in keep:
            sld_ids.remove(sid)
    for n in order:
        sld_ids.remove(ids[n - 1])
        sld_ids.append(ids[n - 1])
    prs.save(out)
    print(f"filled {len(keep)} slides -> {out}")


# ---------- check ----------

def cmd_check(template, map_path, content_path, out, report_path=None):
    tmap = yaml.safe_load(Path(map_path).read_text())
    content = yaml.safe_load(Path(content_path).read_text())
    tpl = Presentation(template)
    filled = Presentation(out)
    issues = []
    keep_keys = set()
    for s in tmap["slides"]:
        for b in s["boxes"]:
            if b.get("keep") or (b.get("chrome") and b["name"] in content.get("chrome", {})):
                keep_keys.add((s["slide"], str(b["key"])))
    banned = [w for w in content.get("leftover_words", [])]

    tpl_styles = {}
    for i, slide in enumerate(tpl.slides, 1):
        for key, shape, tf, cell in boxes(slide):
            tpl_styles[(i, key)] = style(tf)

    samples = {(s["slide"], str(b["key"])): b["sample"].strip() for s in tmap["slides"] for b in s["boxes"]}

    def shape_sample(src, key):  # unchanged template text: its fit is the designer's call
        return samples.get((src, key))

    order = [int(s["slide"]) for s in content["slides"]]
    values = {int(s["slide"]): {str(k): v for k, v in (s.get("boxes") or {}).items()}
              for s in content["slides"]}
    for pos, slide in enumerate(filled.slides):
        src = order[pos]
        seen = set()
        for key, shape, tf, cell in boxes(slide):
            seen.add(key)
            where = f"slide {pos + 1} (template {src}) box {key} '{shape.name}'"
            text = tf.text.strip()
            fixed = (src, key) in keep_keys
            if not fixed and key not in values.get(src, {}) and text:
                issues.append({"kind": "not_filled", "where": where, "text": text})
            for w in banned:
                if w.lower() in text.lower():
                    issues.append({"kind": "leftover", "where": where, "text": text, "word": w})
            before, after = tpl_styles.get((src, key), {}), style(tf)
            drift = {k: (before.get(k), after.get(k)) for k in before if before.get(k) != after.get(k)}
            if drift:
                issues.append({"kind": "style_changed", "where": where, "changed": drift})
            st = after
            if cell is None and st.get("size") and text and not text == shape_sample(src, key):
                w, h = box_size(shape)
                lines = wrap_lines(tf.text, w, st["size"], st.get("bold"))
                need = lines * st["size"] * 1.2
                was = samples.get((src, key), "")
                was_lines = wrap_lines(was, w, st["size"], st.get("bold")) if was else 1
                if need > h * 1.05:
                    issues.append({"kind": "overflow", "where": where, "lines": lines,
                                   "fits_lines": int(h // (st["size"] * 1.2)), "text": text})
                elif was_lines == 1 and lines > 1:
                    issues.append({"kind": "wraps", "where": where, "lines": lines, "text": text})
            if cell is not None and st.get("size") and text:
                col = int(key.split("c")[-1])
                cw = shape.table.columns[col].width / EMU_PER_PT - 14
                lines = wrap_lines(tf.text, cw, st["size"], st.get("bold"))
                if lines > 1:
                    issues.append({"kind": "cell_wraps", "where": where, "lines": lines, "text": text})
        for key in values.get(src, {}):
            if key not in seen and values[src][key] is not None:
                issues.append({"kind": "unknown_key", "where": f"template slide {src}", "key": key})
    report = {"deck": str(out), "slides": len(filled.slides), "issues": issues,
              "ok": not any(i["kind"] in FAILING for i in issues)}
    text = json.dumps(report, indent=2, ensure_ascii=False)
    if report_path:
        Path(report_path).write_text(text + "\n")
    print(text)
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    cmd, *args = sys.argv[1:]
    sys.exit({"map": cmd_map, "fill": cmd_fill, "check": cmd_check}[cmd](*args) or 0)
