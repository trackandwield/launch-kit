"""Build a stand-in branded sample deck for the deck-template test.

No official Atlassian slide template is public, so this makes one in Atlassian's
public colors, built the way designer decks usually are: free text boxes with
default names, shapes, a grouped element, tables and a footer on every slide.
The sample content is about a made-up product so leftover text is easy to spot.

    python tools/make_sample_deck.py companies/atlassian/deck/sample-template.pptx
"""

import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

BLUE = RGBColor(0x00, 0x52, 0xCC)
NAVY = RGBColor(0x17, 0x2B, 0x4D)
SUBTLE = RGBColor(0x5E, 0x6C, 0x84)
LIGHT = RGBColor(0xF4, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x00, 0xB8, 0xD9)
YELLOW = RGBColor(0xFF, 0xAB, 0x00)
PURPLE = RGBColor(0x65, 0x54, 0xC0)
GREEN = RGBColor(0x36, 0xB3, 0x7E)
FONT = "Arial"
FOOTER = "Sample template for testing. Unofficial, not affiliated with Atlassian."

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, paras, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if bullets:
            p.space_after = Pt(10)
        r = p.add_run()
        r.text = ("•  " + line) if bullets else line
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
        r.font.color.rgb = color
    return box


def chrome(slide, n, dark=False):
    """Logo mark, footer and page number, as a master would place them."""
    fg = WHITE if dark else SUBTLE
    mark = rect(slide, Inches(0.6), H - Inches(0.62), Inches(0.28), Inches(0.28),
                WHITE if dark else BLUE, MSO_SHAPE.ISOSCELES_TRIANGLE)
    mark.name = "Logo mark"
    text(slide, Inches(1.0), H - Inches(0.66), Inches(3), Inches(0.36), "Sample Co", 12,
         fg, True).name = "Logo word"
    text(slide, Inches(4.2), H - Inches(0.62), Inches(7.2), Inches(0.3), FOOTER, 9, fg,
         align=PP_ALIGN.CENTER).name = "Footer"
    text(slide, W - Inches(1.2), H - Inches(0.62), Inches(0.6), Inches(0.3), str(n), 10, fg,
         align=PP_ALIGN.RIGHT).name = "Slide number"


def heading(slide, kicker, title):
    rect(slide, Inches(0.6), Inches(0.55), Inches(0.5), Inches(0.08), BLUE)
    text(slide, Inches(0.6), Inches(0.7), Inches(8), Inches(0.4), kicker.upper(), 12, BLUE, True)
    text(slide, Inches(0.6), Inches(1.05), Inches(11.5), Inches(1.0), title, 32, NAVY, True)


def table(slide, x, y, w, rows, col_w=None, first_col_bold=True):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(0.5) * len(rows))
    t = shape.table
    for c in range(len(rows[0])):
        if col_w:
            t.columns[c].width = Emu(int(w * col_w[c]))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.15)
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = val
            run.font.name, run.font.size = FONT, Pt(14)
            run.font.bold = r == 0 or (c == 0 and first_col_bold)
            run.font.color.rgb = WHITE if r == 0 else NAVY
    return shape


# 1. Title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, BLUE)
rect(s, Inches(8.6), Inches(-1.2), Inches(6.2), Inches(6.2), TEAL, MSO_SHAPE.OVAL)
rect(s, Inches(10.4), Inches(3.6), Inches(4.2), Inches(4.2), YELLOW, MSO_SHAPE.OVAL)
text(s, Inches(0.8), Inches(1.6), Inches(3.5), Inches(0.4), "PRODUCT LAUNCH", 14, WHITE, True)
text(s, Inches(0.8), Inches(2.1), Inches(8), Inches(2.0), "Introducing Pipeline Pulse", 54, WHITE, True)
text(s, Inches(0.8), Inches(4.0), Inches(7.2), Inches(1.0),
     "Live health checks for every data pipeline, right where your team already works.", 22, WHITE)
text(s, Inches(0.8), Inches(5.3), Inches(6), Inches(0.4), "Jamie Rivera  ·  Product Marketing  ·  March 2026",
     14, WHITE)
chrome(s, 1, dark=True)

# 2. Agenda
s = prs.slides.add_slide(BLANK)
heading(s, "Agenda", "What we'll cover today")
items = ["The problem with silent pipeline failures", "Meet Pipeline Pulse", "How it works",
         "What customers are saying", "Plans and pricing", "Next steps"]
for i, item in enumerate(items):
    y = Inches(2.3) + Inches(0.68) * i
    num = rect(s, Inches(0.6), y, Inches(0.5), Inches(0.5), BLUE, MSO_SHAPE.OVAL)
    num.text_frame.word_wrap = False
    num.text_frame.margin_left = num.text_frame.margin_right = 0
    num.text_frame.text = f"{i + 1:02d}"
    r = num.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(12), True, WHITE
    num.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, Inches(1.35), y + Inches(0.02), Inches(9), Inches(0.5), item, 20, NAVY,
         anchor=MSO_ANCHOR.MIDDLE)
chrome(s, 2)

# 3. Problem
s = prs.slides.add_slide(BLANK)
heading(s, "The problem", "Broken pipelines are found by customers, not teams")
card = rect(s, Inches(0.6), Inches(2.4), Inches(4.4), Inches(3.9), NAVY, MSO_SHAPE.ROUNDED_RECTANGLE)
card.adjustments[0] = 0.06
text(s, Inches(0.9), Inches(2.8), Inches(3.9), Inches(1.6), "6 hours", 60, YELLOW, True)
text(s, Inches(0.9), Inches(4.3), Inches(3.8), Inches(1.6),
     "Average time before a failed nightly job is noticed by the data team.", 18, WHITE)
text(s, Inches(5.6), Inches(2.5), Inches(6.9), Inches(3.8), [
    "Alerts fire in five different tools, so nobody owns them",
    "Dashboards quietly show stale numbers to leadership",
    "Engineers spend Monday mornings tracing what broke on Friday",
    "Every fix starts with the same question: who changed what?",
], 20, NAVY, bullets=True)
chrome(s, 3)

# 4. Solution: three pillars
s = prs.slides.add_slide(BLANK)
heading(s, "The solution", "One place to see, fix and prevent pipeline issues")
pillars = [("See it first", "Health checks run after every job and flag problems in seconds.", BLUE),
           ("Fix it fast", "Every alert links to the run, the change and the owner.", PURPLE),
           ("Stop repeats", "Weekly reports show which pipelines break most and why.", GREEN)]
for i, (t, body, col) in enumerate(pillars):
    x = Inches(0.6) + Inches(4.1) * i
    rect(s, x, Inches(2.4), Inches(3.8), Inches(3.9), LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.05
    rect(s, x + Inches(0.35), Inches(2.75), Inches(0.75), Inches(0.75), col, MSO_SHAPE.OVAL)
    text(s, x + Inches(0.35), Inches(3.75), Inches(3.2), Inches(0.6), t, 24, NAVY, True)
    text(s, x + Inches(0.35), Inches(4.4), Inches(3.2), Inches(1.7), body, 16, SUBTLE)
chrome(s, 4)

# 5. How it works: grouped steps
s = prs.slides.add_slide(BLANK)
heading(s, "How it works", "Up and running in three steps")
grp = s.shapes.add_group_shape()
grp.name = "Steps"
steps = [("Connect", "Link your scheduler in two clicks. No agents to install."),
         ("Choose checks", "Pick from 40 ready-made checks or write your own in SQL."),
         ("Get alerts", "Issues land in your team's chat with the owner tagged.")]
for i, (t, body) in enumerate(steps):
    x = Inches(0.6) + Inches(4.1) * i
    chev = grp.shapes.add_shape(MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, x, Inches(2.5),
                                Inches(3.9), Inches(0.9))
    chev.fill.solid()
    chev.fill.fore_color.rgb = [BLUE, TEAL, NAVY][i]
    chev.line.fill.background()
    chev.text_frame.text = f"Step {i + 1}"
    r = chev.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(16), True, WHITE
    tb = grp.shapes.add_textbox(x, Inches(3.7), Inches(3.6), Inches(0.6))
    tb.text_frame.text = t
    r = tb.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(22), True, NAVY
    tb = grp.shapes.add_textbox(x, Inches(4.35), Inches(3.6), Inches(1.6))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = body
    r = tb.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(16), SUBTLE
chrome(s, 5)

# 6. Proof: quote and stats
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, Inches(7.4), H, NAVY)
text(s, Inches(0.8), Inches(0.9), Inches(1.5), Inches(1.4), "“", 120, TEAL, True)
text(s, Inches(0.8), Inches(2.2), Inches(6.0), Inches(2.6),
     "We used to hear about broken reports from the CFO. Now we fix them before anyone opens a dashboard.",
     26, WHITE, True)
text(s, Inches(0.8), Inches(5.1), Inches(6), Inches(0.8), ["Priya Shah", "Head of Data, Northwind Outfitters"],
     16, WHITE)
for i, (num, label) in enumerate([("83%", "fewer stale dashboards"), ("4x", "faster time to fix"),
                                  ("12 min", "median time to detect")]):
    y = Inches(1.0) + Inches(1.8) * i
    text(s, Inches(8.2), y, Inches(4.5), Inches(0.9), num, 48, BLUE, True)
    text(s, Inches(8.2), y + Inches(0.85), Inches(4.5), Inches(0.5), label, 18, SUBTLE)
chrome(s, 6)

# 7. Pricing table
s = prs.slides.add_slide(BLANK)
heading(s, "Plans", "Simple pricing that grows with you")
table(s, Inches(0.6), Inches(2.4), Inches(12.1), [
    ["", "Free", "Standard", "Premium"],
    ["Price", "$0", "$12 per user/month", "$20 per user/month"],
    ["Pipelines", "Up to 5", "Unlimited", "Unlimited"],
    ["Checks", "Ready-made only", "Ready-made and custom", "Ready-made and custom"],
    ["Support", "Community", "Business hours", "24/7 with 1-hour response"],
], col_w=[0.22, 0.22, 0.28, 0.28])
text(s, Inches(0.6), Inches(5.4), Inches(11), Inches(0.4), "Save 15% with annual billing.", 14, SUBTLE)
chrome(s, 7)

# 8. Comparison
s = prs.slides.add_slide(BLANK)
heading(s, "Why us", "How Pipeline Pulse compares")
table(s, Inches(0.6), Inches(2.4), Inches(12.1), [
    ["", "Pipeline Pulse", "DIY scripts", "Legacy monitoring"],
    ["Setup time", "Minutes", "Weeks", "Days"],
    ["Knows who owns each job", "Yes", "No", "Partly"],
    ["Alerts in team chat", "Yes", "If you build it", "Email only"],
    ["Cost to maintain", "Included", "Engineer time", "Annual contract"],
], col_w=[0.31, 0.23, 0.23, 0.23])
chrome(s, 8)

# 9. Closing / CTA
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, BLUE)
rect(s, Inches(-1.5), Inches(3.8), Inches(5), Inches(5), PURPLE, MSO_SHAPE.OVAL)
text(s, Inches(4.0), Inches(2.0), Inches(8.5), Inches(1.2), "Try Pipeline Pulse free", 48, WHITE, True)
text(s, Inches(4.0), Inches(3.3), Inches(8.3), Inches(1.0),
     "Connect your first pipeline in under five minutes. No credit card needed.", 22, WHITE)
btn = rect(s, Inches(4.0), Inches(4.6), Inches(3.4), Inches(0.75), YELLOW, MSO_SHAPE.ROUNDED_RECTANGLE)
btn.text_frame.text = "samplecompany.com/pulse"
r = btn.text_frame.paragraphs[0].runs[0]
r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(16), True, NAVY
chrome(s, 9, dark=True)

out = sys.argv[1] if len(sys.argv) > 1 else "sample-template.pptx"
prs.save(out)
print("wrote", out)
